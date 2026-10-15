"""Experimental: HTML deck -> Keynote with movable pictures as well as editable text.

A PDF only says "draw this image here"; the HTML a deck was made from says
which element is a photo, how it is cropped, and what is layered over it. So:

1. Chrome (headless, via Playwright) lays out each slide on its own page at
   the deck's size, with the deck's fonts.
2. Every element is listed in paint order and sorted into
   - pictures: <img> and CSS background images. They become real Keynote
     pictures from the original files, with the crop the browser applies
     (object-fit / object-position / background-size, overflow clipping),
     opacity and rounded corners. Pictures with effects Keynote can't
     reproduce (filters, masks, rotation, blend modes) are captured on
     their own instead, still as a separate, movable picture.
   - decorations: backgrounds, cards, gradients, borders, SVG. Each run of
     decorations between two pictures is captured as one transparent layer,
     trimmed to its content, so overlays stay above the photos they cover.
3. Chrome also prints each slide to PDF, and the regular text pipeline turns
   that into editable text boxes on top.

Slides are the deck's <section> elements (Claude Design exports, reveal-style
decks); a page without sections becomes a single slide.
"""
import base64
import json
import re
import tempfile
import zipfile
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote, urlparse

import pymupdf
from PIL import Image

from pptx.oxml.ns import qn

from . import converter
from . import report as reporting

PX_TO_PT = 0.75
SCALE = 2           # screenshot pixels per CSS pixel
MAX_IMAGE_SIDE = 3200

# Everything the page knows about each element, in paint order.
COLLECT_JS = r"""
(root) => {
  const items = [];
  const all = [root, ...root.querySelectorAll('*')];
  const rootRect = root.getBoundingClientRect();
  const R = r => ({x: r.left - rootRect.left, y: r.top - rootRect.top, w: r.width, h: r.height});
  let order = 0;
  for (const el of all) {
    el.setAttribute('data-p2k', String(order++));
  }
  const inSvg = el => el.closest('svg') && el.tagName.toLowerCase() !== 'svg';
  const effZ = el => {
    for (let e = el; e && e !== root.parentElement; e = e.parentElement) {
      const cs = getComputedStyle(e);
      if (cs.position !== 'static' && cs.zIndex !== 'auto') return parseInt(cs.zIndex) || 0;
    }
    return 0;
  };
  const complex = el => {   // effects a Keynote picture can't reproduce
    for (let e = el; e && e !== root.parentElement; e = e.parentElement) {
      const cs = getComputedStyle(e);
      if (cs.filter !== 'none' || cs.mixBlendMode !== 'normal' || cs.clipPath !== 'none'
          || (cs.maskImage && cs.maskImage !== 'none') || (cs.webkitMaskImage && cs.webkitMaskImage !== 'none'))
        return true;
      if (cs.transform !== 'none') {
        const m = new DOMMatrix(cs.transform);
        if (Math.abs(m.b) > 1e-3 || Math.abs(m.c) > 1e-3) return true;   // rotated / skewed
      }
    }
    return false;
  };
  const opacity = el => {
    let o = 1;
    for (let e = el; e && e !== root.parentElement; e = e.parentElement) o *= parseFloat(getComputedStyle(e).opacity);
    return o;
  };
  const clipBox = el => {   // visible area after overflow clipping by ancestors
    let c = {l: -1e9, t: -1e9, r: 1e9, b: 1e9}, radius = null;
    for (let e = el.parentElement; e && e !== root.parentElement; e = e.parentElement) {
      const cs = getComputedStyle(e);
      if (cs.overflow !== 'visible' || cs.overflowX !== 'visible' || cs.overflowY !== 'visible') {
        const r = e.getBoundingClientRect();
        if (r.width * r.height < (c.r - c.l) * (c.b - c.t)) radius = parseFloat(cs.borderTopLeftRadius) || 0;
        c = {l: Math.max(c.l, r.left), t: Math.max(c.t, r.top), r: Math.min(c.r, r.right), b: Math.min(c.b, r.bottom)};
      }
    }
    return {x: c.l - rootRect.left, y: c.t - rootRect.top, w: c.r - c.l, h: c.b - c.t, radius};
  };
  for (const el of all) {
    if (inSvg(el)) continue;
    const cs = getComputedStyle(el);
    if (cs.display === 'none' || cs.visibility === 'hidden') continue;
    const r = el.getBoundingClientRect();
    if (r.width < 0.5 || r.height < 0.5) continue;
    const tag = el.tagName.toLowerCase();
    const base = {id: el.getAttribute('data-p2k'), tag, rect: R(r), z: effZ(el), order: +el.getAttribute('data-p2k')};
    const bg = cs.backgroundImage;
    const urls = bg && bg !== 'none' ? [...bg.matchAll(/url\("?([^")]*)"?\)/g)].map(m => m[1]) : [];
    if (tag === 'img' && el.currentSrc && el.naturalWidth) {
      items.push({...base, kind: 'image', src: el.currentSrc, natural: [el.naturalWidth, el.naturalHeight],
        fit: cs.objectFit, pos: cs.objectPosition, opacity: opacity(el), complex: complex(el),
        radius: parseFloat(cs.borderTopLeftRadius) || 0, clip: clipBox(el),
        content: R(new DOMRect(r.left + parseFloat(cs.borderLeftWidth) + parseFloat(cs.paddingLeft),
                               r.top + parseFloat(cs.borderTopWidth) + parseFloat(cs.paddingTop),
                               el.clientWidth - parseFloat(cs.paddingLeft) - parseFloat(cs.paddingRight),
                               el.clientHeight - parseFloat(cs.paddingTop) - parseFloat(cs.paddingBottom)))});
      // the img's own box (border, background) still paints as decoration
    }
    if (urls.length === 1 && !/gradient\(/.test(bg)) {
      items.push({...base, kind: 'bgimage', src: urls[0], size: cs.backgroundSize, pos: cs.backgroundPosition,
        repeat: cs.backgroundRepeat, opacity: opacity(el), complex: complex(el),
        radius: parseFloat(cs.borderTopLeftRadius) || 0, clip: clipBox(el)});
    }
    const pseudo = ['::before', '::after'].some(p => getComputedStyle(el, p).content !== 'none'
                                                   && getComputedStyle(el, p).content !== 'normal');
    const paints = tag !== 'img' && (el === root || tag === 'svg' || tag === 'canvas' || tag === 'video' || tag === 'hr'
      || tag === 'iframe' || pseudo
      || !/rgba\(0, 0, 0, 0\)|transparent/.test(cs.backgroundColor)
      || (bg !== 'none' && (urls.length !== 1 || /gradient\(/.test(bg)))
      || ['Top', 'Right', 'Bottom', 'Left'].some(s => parseFloat(cs['border' + s + 'Width']) > 0
                                                && cs['border' + s + 'Style'] !== 'none')
      || cs.boxShadow !== 'none' || (cs.outlineStyle !== 'none' && parseFloat(cs.outlineWidth) > 0));
    if (paints) items.push({...base, kind: 'decoration'});
  }
  items.sort((a, b) => a.z - b.z || a.order - b.order);
  return items;
}
"""

BASE_CSS = """
*, *::before, *::after { animation: none !important; transition: none !important; }
html, body { margin: 0 !important; padding: 0 !important; background: transparent !important; overflow: hidden !important; }
"""
HIDE_TEXT_CSS = """
[data-p2k-root], [data-p2k-root] * { -webkit-text-fill-color: transparent !important; text-shadow: none !important;
  -webkit-text-stroke: 0 !important; caret-color: transparent !important; text-decoration-color: transparent !important; }
[data-p2k-root] svg text { fill: transparent !important; stroke: none !important; }
"""


NOT_THE_DECK = re.compile(r"print|backup|parked|export|copy|old|draft|thumbnail|\bref\b", re.I)


def _slide_count(html_file):
    try:
        return len(re.findall(r"<section\b", html_file.read_text(encoding="utf-8", errors="ignore")))
    except OSError:
        return 0


def _version(name):
    m = re.search(r"\bv(\d+)\b", name, re.I)
    return int(m.group(1)) if m else 0


def list_decks(path):
    """HTML decks in an export folder, most likely first, as (relative path,
    slide count). Exports often hold old versions, print copies and reference
    pages beside the deck; those rank last."""
    root = Path(path)
    found = []
    for f in root.rglob("*.htm*"):
        rel = f.relative_to(root)
        if f.name.startswith(".") or f.suffix.lower() not in (".html", ".htm"):
            continue
        n = _slide_count(f)
        if n:
            found.append((rel, n))
    def rank(item):
        rel, n = item
        sidelined = bool(NOT_THE_DECK.search(str(rel))) or len(rel.parts) > min(len(r.parts) for r, _ in found)
        return (sidelined, -_version(rel.stem), -n, str(rel))
    return [(str(r), n) for r, n in sorted(found, key=rank)]


def _unpack(path, workdir):
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():  # no absolute paths or ".." escapes
            if name.startswith("/") or ".." in Path(name).parts:
                raise ValueError(f"unsafe path in zip: {name}")
        z.extractall(workdir)
    return Path(workdir)


def decks_in(path):
    """Decks inside a .zip export or folder (no files are left behind)."""
    path = Path(path)
    if path.suffix.lower() == ".zip":
        with tempfile.TemporaryDirectory() as tmp:
            return list_decks(_unpack(path, tmp))
    if path.is_dir():
        return list_decks(path)
    return [(path.name, _slide_count(path))]


def _find_html(path, workdir, deck=None):
    """The deck's HTML file, from a .zip export, an unzipped export folder, or
    the .html itself. Its assets (images, styles, fonts) must sit beside it.
    `deck` picks one by (part of) its name when the export holds several."""
    path = Path(path)
    if path.suffix.lower() == ".zip":
        path = _unpack(path, workdir)
    if path.is_dir():
        decks = list_decks(path)
        if not decks:
            raise ValueError(f"no HTML deck (a page with <section> slides) in {path.name}")
        if deck:
            chosen = [d for d, _ in decks if d == deck] or [d for d, _ in decks if deck.lower() in d.lower()]
            if not chosen:
                raise ValueError(f"no deck matching {deck!r}; found: " + "; ".join(d for d, _ in decks))
            return (path / chosen[0]).resolve()
        return (path / decks[0][0]).resolve()
    if path.suffix.lower() not in (".html", ".htm"):
        raise ValueError("expected a .zip export, an export folder, or an .html file")
    return path.resolve()


def _slide_size(page):
    size = page.evaluate("""() => {
      const x = document.querySelector('[width][height]:not(img):not(svg):not(canvas):not(video):not(iframe)');
      return x ? [parseInt(x.getAttribute('width')), parseInt(x.getAttribute('height'))] : null }""")
    return tuple(size) if size and all(size) else (1920, 1080)


def _head_markup(page):
    """Stylesheets and fonts of the deck, to rebuild each slide on its own page."""
    return page.evaluate("""() => [...document.querySelectorAll('link[rel=stylesheet], style')]
                              .map(e => e.outerHTML).join('\\n')""")


def _render_layer(page, visible, hide_bg_images):
    """Screenshot with only `visible` elements painted (no text, no pictures)."""
    rules = ["[data-p2k-root], [data-p2k-root] *:not(svg *) { visibility: hidden !important; }"]
    if visible:
        # :not(#p2k) adds an id's worth of specificity, so these beat the rule above
        sel = ", ".join(f'[data-p2k="{i}"]:not(#p2k)' for i in visible)
        rules.append(f"{sel} {{ visibility: visible !important; }}")
        # An <svg>'s children inherit its visibility, so they need no rule of their
        # own; forcing them visible would undo parts the deck hides on purpose
        # (build-up stages like `section.v1 .s2 {visibility: hidden}`).
    if hide_bg_images:
        rules.append(", ".join(f'[data-p2k="{i}"]:not(#p2k)' for i in hide_bg_images)
                     + " { background-image: none !important; }")
    page.evaluate("css => { document.getElementById('p2k-layer').textContent = css }", "\n".join(rules))
    png = page.screenshot(omit_background=True, type="png")
    return Image.open(BytesIO(png)).convert("RGBA")


def _local_file(src, base_dir):
    u = urlparse(src)
    if u.scheme == "file":
        return Path(unquote(u.path))
    if u.scheme in ("", None):
        return (base_dir / unquote(src)).resolve()
    if u.scheme == "data":
        return None
    return None


def _image_bytes(src, base_dir):
    if src.startswith("data:"):
        return base64.b64decode(src.split(",", 1)[1]) if ";base64," in src[:100] else None
    f = _local_file(src, base_dir)
    return f.read_bytes() if f and f.exists() else None


def _parse_pos(value, free):
    """CSS position component ('50%', '12px', 'left') -> offset in px."""
    value = value.strip()
    keywords = {"left": 0, "top": 0, "center": 50, "right": 100, "bottom": 100}
    if value in keywords:
        return free * keywords[value] / 100
    if value.endswith("%"):
        return free * float(value[:-1]) / 100
    if value.endswith("px"):
        return float(value[:-2])
    return free / 2


def _displayed_rect(item):
    """Where the whole image is drawn (x, y, w, h), before clipping."""
    iw, ih = item["natural"]
    if item["kind"] == "image":
        box = item["content"]
        fit = item["fit"]
        pos = item["pos"].split()
    else:
        box = item["rect"]
        size = item["size"]
        fit = "cover" if size == "cover" else "contain" if size == "contain" else "explicit"
        pos = item["pos"].split()
        if fit == "explicit":
            parts = size.split()
            def dim(v, total, other_auto_ratio=None):
                if v.endswith("%"):
                    return total * float(v[:-1]) / 100
                if v.endswith("px"):
                    return float(v[:-2])
                return None
            w = dim(parts[0], box["w"]) if parts else None
            h = dim(parts[1], box["h"]) if len(parts) > 1 else None
            if w is None and h is None:
                w, h = iw, ih
            elif w is None:
                w = h * iw / ih
            elif h is None:
                h = w * ih / iw
            px, py = (pos + ["50%", "50%"])[:2]
            return (box["x"] + _parse_pos(px, box["w"] - w), box["y"] + _parse_pos(py, box["h"] - h), w, h)
    bw, bh = box["w"], box["h"]
    if fit == "cover":
        k = max(bw / iw, bh / ih)
    elif fit == "contain":
        k = min(bw / iw, bh / ih)
    elif fit == "none":
        k = 1
    elif fit == "scale-down":
        k = min(1, bw / iw, bh / ih)
    else:  # fill
        return (box["x"], box["y"], bw, bh)
    w, h = iw * k, ih * k
    px, py = (pos + ["50%", "50%"])[:2]
    return (box["x"] + _parse_pos(px, bw - w), box["y"] + _parse_pos(py, bh - h), w, h)


PPTX_FORMATS = {"BMP", "GIF", "JPEG", "PNG", "TIFF"}


def _keynote_ready(img, data):
    """Bytes of a format the .pptx can carry. AVIF, WebP, HEIC and friends
    (often hiding behind a .png or .jpg name) are converted to PNG or JPEG."""
    if img.format in PPTX_FORMATS:
        return data
    buf = BytesIO()
    if img.mode in ("RGBA", "LA", "P", "PA") or "transparency" in img.info:
        img.convert("RGBA").save(buf, "PNG", optimize=True)
    else:
        img.convert("RGB").save(buf, "JPEG", quality=92)
    return buf.getvalue()


def _decodable(src, base_dir):
    """Can the original bytes be read here (else the browser's rendering is used)."""
    data = _image_bytes(src, base_dir)
    if not data:
        return False
    try:
        Image.open(BytesIO(data)).size
        return True
    except Exception:
        return False


def _photo_name(item):
    name = Path(unquote(urlparse(item["src"]).path)).name if not item["src"].startswith("data:") else "image"
    return f"Photo · {name}"


def _add_picture(shapes, item, base_dir, layer_png=None):
    """A real picture: original file, cropped like the browser crops it.
    Returns (shape, visible rect in CSS px) or None."""
    k = PX_TO_PT
    if layer_png is not None:  # captured on its own (effects Keynote can't do)
        bbox = layer_png.getbbox()
        if not bbox:
            return None
        pic = _add_png(shapes, bbox, layer_png.crop(bbox), _photo_name(item))
        return pic, tuple(v / SCALE for v in bbox)
    data = _image_bytes(item["src"], base_dir)
    if not data:
        return None
    try:
        img = Image.open(BytesIO(data))
        data = _keynote_ready(img, data)
    except Exception:
        return None
    item.setdefault("natural", list(img.size))
    dx, dy, dw, dh = _displayed_rect(item)
    # visible part: displayed image ∩ its box ∩ ancestors' overflow clip
    box = item["content"] if item["kind"] == "image" else item["rect"]
    c = item["clip"]
    vx0, vy0 = max(dx, box["x"], c["x"]), max(dy, box["y"], c["y"])
    vx1 = min(dx + dw, box["x"] + box["w"], c["x"] + c["w"])
    vy1 = min(dy + dh, box["y"] + box["h"], c["y"] + c["h"])
    if vx1 - vx0 < 0.5 or vy1 - vy0 < 0.5:
        return None

    # shrink huge originals to what the slide can show at retina density
    need = max((vx1 - vx0) / dw * img.width, (vy1 - vy0) / dh * img.height)
    show_px = max(vx1 - vx0, vy1 - vy0) * SCALE
    if max(img.size) > MAX_IMAGE_SIDE and need > show_px * 1.5:
        f = max(show_px * 1.5 / need, MAX_IMAGE_SIDE / max(img.size))
        img = img.resize((max(1, round(img.width * f)), max(1, round(img.height * f))), Image.LANCZOS)
        buf = BytesIO()
        if img.mode in ("RGBA", "LA", "P"):
            img.save(buf, "PNG", optimize=True)
        else:
            img.convert("RGB").save(buf, "JPEG", quality=90)
        data = buf.getvalue()

    pic = shapes.add_picture(BytesIO(data), converter._emu(vx0 * k), converter._emu(vy0 * k),
                             converter._emu((vx1 - vx0) * k), converter._emu((vy1 - vy0) * k))
    pic.name = _photo_name(item)
    pic.crop_left = (vx0 - dx) / dw
    pic.crop_top = (vy0 - dy) / dh
    pic.crop_right = (dx + dw - vx1) / dw
    pic.crop_bottom = (dy + dh - vy1) / dh

    if item["opacity"] < 0.999:
        blip = pic._element.blipFill.find(qn("a:blip"))
        blip.append(blip.makeelement(qn("a:alphaModFix"), {"amt": str(round(item["opacity"] * 100000))}))
    radius = item["radius"] or (c.get("radius") or 0)
    short = min(vx1 - vx0, vy1 - vy0)
    if radius > 0.5 and short > 0:
        # rounded-rectangle mask; adj is the radius as a share of the short side (max 50000)
        geom = pic._element.spPr.find(qn("a:prstGeom"))
        geom.set("prst", "roundRect")
        av = geom.find(qn("a:avLst"))
        if av is None:
            av = geom.makeelement(qn("a:avLst"), {})
            geom.append(av)
        av.append(av.makeelement(qn("a:gd"), {"name": "adj", "fmla": f"val {round(min(0.5, radius / short) * 100000)}"}))
    return pic, (vx0, vy0, vx1, vy1)


MAX_PIECES = 80


def _add_png(shapes, bbox, img, name):
    """Place an RGBA crop whose top-left sits at screenshot pixel bbox[:2]."""
    buf = BytesIO()
    img.save(buf, "PNG", optimize=True)
    k = PX_TO_PT / SCALE
    pic = shapes.add_picture(BytesIO(buf.getvalue()), converter._emu(bbox[0] * k), converter._emu(bbox[1] * k),
                             converter._emu(img.width * k), converter._emu(img.height * k))
    pic.name = name
    return pic


def _pieces(png):
    """Split a decoration layer into its separate islands of paint (a card, an
    arrow, a badge, an overlay), each as (bbox, cropped RGBA), so they become
    separate objects. Parts closer than ~2 CSS px stay together."""
    import numpy as np
    from scipy import ndimage

    bbox = png.getbbox()
    if not bbox:
        return []
    arr = np.asarray(png)
    ink = arr[:, :, 3] > 8
    labels, n = ndimage.label(ndimage.binary_dilation(ink, iterations=2 * SCALE))
    if n <= 1 or n > MAX_PIECES:
        return [(bbox, png.crop(bbox))]
    out = []
    for k, sl in enumerate(ndimage.find_objects(labels), start=1):
        ys, xs = sl
        piece = arr[ys, xs].copy()
        piece[labels[ys, xs] != k, 3] = 0  # parts of other islands inside this box
        img = Image.fromarray(piece, "RGBA")
        inner = img.getbbox()
        if not inner:
            continue
        x0, y0 = xs.start + inner[0], ys.start + inner[1]
        out.append(((x0, y0, xs.start + inner[2], ys.start + inner[3]), img.crop(inner)))
    return out


def _piece_name(bbox, width, height):
    w, h = (bbox[2] - bbox[0]) / SCALE, (bbox[3] - bbox[1]) / SCALE
    if w >= 0.95 * width and h >= 0.95 * height:
        return "Background"
    return "Shape"


def _inside(bbox, rect, slack=2):
    """Screenshot-pixel bbox lies within a CSS-px rect."""
    x0, y0, x1, y1 = (v / SCALE for v in bbox)
    return x0 >= rect[0] - slack and y0 >= rect[1] - slack and x1 <= rect[2] + slack and y1 <= rect[3] + slack


def _build_slide(slide, layers, base_dir, width, height, report=None):
    """Add the slide's layers bottom to top. A photo and the overlays that sit
    only on it (a card's gradient, a caption plate) become one group, so they
    move together and the photo stays reachable by double-clicking."""
    def split(layer):
        if len(layer) > 2:  # full-slide paint stays whole
            bbox = layer[1].getbbox()
            return [(bbox, layer[1].crop(bbox))] if bbox else []
        return _pieces(layer[1])

    pieces = [split(l) if l[0] == "layer" else None for l in layers]
    first = True
    for i, layer in enumerate(layers):
        if layer[0] == "layer":
            for bbox, img in pieces[i]:
                name = _piece_name(bbox, width, height)
                if name == "Background" and not first:
                    name = "Overlay"
                first = False
                _add_png(slide.shapes, bbox, img, name)
            continue
        first = False
        nxt = pieces[i + 1] if i + 1 < len(layers) and layers[i + 1][0] == "layer" else []
        placed = _add_picture(slide.shapes, layer[1], base_dir, layer[2])
        if not placed:
            if report is not None and layer[2] is None:
                report["warnings"].append(f"image not found: {layer[1]['src'][:120]}")
            continue
        pic, rect = placed
        on_top = [pc for pc in nxt if _inside(pc[0], rect)]
        if on_top:
            group = slide.shapes.add_group_shape()
            group.name = pic.name.replace("Photo", "Photo group")
            group.shapes._spTree.append(pic._element)  # move the photo into the group
            for bbox, img in on_top:
                _add_png(group.shapes, bbox, img, "Overlay")
            group.shapes._recalculate_extents()
            for pc in on_top:
                nxt.remove(pc)


def html_to_pptx(html_path, pptx_path, pdf_out, progress=None, report=None, deck=None):
    """Build the .pptx from an HTML deck (or a zip of one). Also writes the
    deck as a PDF to `pdf_out` (used for the text and the accuracy check)."""
    from playwright.sync_api import sync_playwright

    with tempfile.TemporaryDirectory() as tmp, sync_playwright() as pw:
        html = _find_html(html_path, tmp, deck)
        if report is not None:
            report["deck"] = html.name
        base_dir = html.parent
        browser = pw.chromium.launch(channel="chrome")
        try:
            page = browser.new_page(viewport={"width": 1920, "height": 1080})
            page.goto(html.as_uri(), wait_until="load")
            page.wait_for_timeout(1500)
            width, height = _slide_size(page)
            head = _head_markup(page)
            sections = page.evaluate("""() => {
              const s = [...document.querySelectorAll('section')].filter(e => !e.parentElement.closest('section'));
              return s.length ? s.map(e => e.outerHTML) : [document.body.innerHTML] }""")

            slide_page = browser.new_page(viewport={"width": width, "height": height}, device_scale_factor=SCALE)
            slide_page.emulate_media(media="screen")
            pdf = pymupdf.open()
            plans = []
            slide_file = base_dir / f".pdf2key-slide-{next(tempfile._get_candidate_names())}.html"
            for i, markup in enumerate(sections):
                if progress:
                    progress(f"Laying out slide {i + 1} of {len(sections)}", 0.25 * i / len(sections))
                doc_html = f"""<!doctype html><html><head><meta charset="utf-8">
                    {head}<style>{BASE_CSS}</style><style id="p2k-layer"></style></head><body>{markup}</body></html>"""
                # a file next to the deck, so relative images and fonts load (file:// pages
                # can't be reached from about:blank)
                slide_file.write_text(doc_html, encoding="utf-8")
                slide_page.goto(slide_file.as_uri(), wait_until="load")
                slide_page.evaluate(f"""() => {{
                    const s = document.body.firstElementChild;
                    s.setAttribute('data-p2k-root', '');
                    Object.assign(s.style, {{position: 'absolute', left: '0', top: '0', width: '{width}px',
                        height: '{height}px', opacity: '1', transform: 'none', visibility: 'visible',
                        display: getComputedStyle(s).display === 'none' ? 'block' : getComputedStyle(s).display}});
                }}""")
                slide_page.evaluate("() => document.fonts.ready")
                slide_page.wait_for_timeout(150)
                pdf_bytes = slide_page.pdf(width=f"{width}px", height=f"{height}px", print_background=True,
                                           page_ranges="1", margin={"top": "0", "right": "0", "bottom": "0", "left": "0"})
                pdf.insert_pdf(pymupdf.open("pdf", pdf_bytes))

                items = slide_page.evaluate(COLLECT_JS, slide_page.query_selector("[data-p2k-root]"))
                slide_page.add_style_tag(content=HIDE_TEXT_CSS)
                bg_image_ids = [it["id"] for it in items if it["kind"] == "bgimage"]
                layers, run = [], []

                def flush():
                    # Full-slide paint (the background, a tint over everything) gets its
                    # own layer, in place; otherwise it would fuse every shape above it
                    # into one picture when the layer is split into pieces.
                    sub = []
                    for it in run:
                        r = it["rect"]
                        if r["w"] * r["h"] >= 0.9 * width * height:
                            if sub:
                                layers.append(("layer", _render_layer(slide_page, sub, bg_image_ids)))
                                sub = []
                            layers.append(("layer", _render_layer(slide_page, [it["id"]], bg_image_ids), "full"))
                        else:
                            sub.append(it["id"])
                    if sub:
                        layers.append(("layer", _render_layer(slide_page, sub, bg_image_ids)))
                    run.clear()

                for it in items:
                    if it["kind"] == "decoration":
                        run.append(it)
                        continue
                    if run:
                        flush()
                    # effects Keynote can't do, or a file we can't read here (SVG, remote
                    # URL, corrupt): use the browser's own rendering of just this picture
                    if it["complex"] or not _decodable(it["src"], base_dir):
                        layers.append(("picture", it, _render_layer(slide_page, [it["id"]], [])))
                    else:
                        layers.append(("picture", it, None))
                if run:
                    flush()
                plans.append(layers)
            pdf.save(pdf_out)
        finally:
            browser.close()
            if "slide_file" in locals():
                slide_file.unlink(missing_ok=True)

        doc = converter.open_pdf(pdf_out)
        pages_blocks = converter.analyze_text(doc, None, report)
        w_pt, h_pt = width * PX_TO_PT, height * PX_TO_PT
        prs = converter.new_presentation(w_pt, h_pt)
        blank = prs.slide_layouts[6]
        for i, layers in enumerate(plans):
            if progress:
                progress(f"Building slide {i + 1} of {len(plans)}", 0.4 + 0.3 * i / len(plans))
            slide = prs.slides.add_slide(blank)
            _build_slide(slide, layers, base_dir, width, height, report)
            converter.add_text(slide, pages_blocks[i], doc[i].rect, w_pt, h_pt, report)
        reporting.record_objects(report, prs)
        prs.save(pptx_path)


def convert_html(html_path, key_path, verify=True, progress=None, preview_dir=None, report=None, deck=None,
                 pptx_out=None, keynote=True):
    """HTML deck (or zip/folder) -> .key and/or .pptx, like converter.convert.
    Pass a dict from report.new() as `report` to receive the status report."""
    import time
    started = time.time()
    key_path = Path(key_path).resolve() if key_path else None
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        pptx, pdf = tmp / "deck.pptx", tmp / "deck.pdf"
        html_to_pptx(html_path, pptx, pdf, progress, report, deck)
        slides = converter._finish(pptx, key_path, pptx_out, keynote, verify, pdf, progress, preview_dir, tmp)
        reporting.record_accuracy(report, slides, time.time() - started)
        return slides

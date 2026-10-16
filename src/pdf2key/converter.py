"""Core PDF -> editable Keynote conversion.

Each PDF page becomes a slide with two layers:
  1. art  - everything except text (photos, shapes, gradients, logos) rendered
            exactly as the PDF draws it, so stacking, cropping, transparency
            and masks are always right. One high-resolution picture.
  2. text - real, editable text boxes with the PDF's line breaks, line spacing,
            letter-spacing, size, color, opacity, weight and rotation.

With separate_images=True the art layer is instead split into individually
movable pictures plus a flattened line-art overlay; easier to rearrange, but
less faithful when the PDF crops, masks, or layers its images.

Pipeline: PDF --PyMuPDF--> .pptx (python-pptx) --Keynote (AppleScript)--> .key
Afterwards every slide is exported from Keynote and compared with the PDF page.
"""
import math
import re
import shutil
import statistics
import subprocess
import sys
import tempfile
import time
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import numpy as np
import pymupdf
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from . import fonts
from . import report as reporting

PT_TO_EMU = 12700
ART_WIDTH_PX = 3840  # art layer resolution (4K across the slide)
FONT_DIRS = [Path("/Library/Fonts"), Path("/System/Library/Fonts"), Path.home() / "Library/Fonts"]
WEIGHT_WORDS = {"thin", "extralight", "ultralight", "light", "medium", "semibold", "demibold",
                "extrabold", "ultrabold", "black", "heavy", "condensed", "narrow"}
BOLD_TOKENS = {"bold", "boldital", "bolditalic", "boldoblique", "bd", "bdit"}


# --------------------------------------------------------------------------- #
# Geometry: map PDF points onto the slide (pages of different sizes are fit).
# --------------------------------------------------------------------------- #
def _emu(v):
    return Emu(int(round(v * PT_TO_EMU)))


class _Fit:
    def __init__(self, rect, slide_w, slide_h):
        self.s = min(slide_w / rect.width, slide_h / rect.height)
        self.ox = (slide_w - rect.width * self.s) / 2
        self.oy = (slide_h - rect.height * self.s) / 2

    def x(self, v):
        return _emu(self.ox + v * self.s)

    def y(self, v):
        return _emu(self.oy + v * self.s)

    def d(self, v):
        return _emu(max(v, 0) * self.s)


def _page_copy(page):
    doc = pymupdf.open()
    doc.insert_pdf(page.parent, from_page=page.number, to_page=page.number)
    return doc


# --------------------------------------------------------------------------- #
# Fonts
# --------------------------------------------------------------------------- #
_font_stems = None


def _norm(name):
    return re.sub(r"[^a-z0-9]", "", name.lower())


def _installed_stems():
    global _font_stems
    if _font_stems is None:
        _font_stems = set()
        for d in FONT_DIRS:
            if d.exists():
                for f in d.rglob("*"):
                    if f.suffix.lower() in (".ttf", ".otf", ".ttc", ".dfont"):
                        _font_stems.add(_norm(re.split(r"[-\[]", f.stem)[0]))
    return _font_stems


def _is_installed(base):
    n = re.sub(r"(mt|ps)$", "", _norm(base))
    stems = _installed_stems()
    return n in stems or (len(n) >= 4 and any(s.startswith(n) for s in stems))


def _anonymous(pdf_name):
    """Fonts without a usable name (e.g. Type3 glyph fonts from Chrome/Skia)."""
    name = pdf_name.split("+")[-1].strip()
    return not name or name.startswith("Type3")


def _font_info(pdf_name, flags, mono=False, weight=None, slanted=False, face=None):
    """Return (family, bold, italic, measuring_font, exact).

    exact=False means the real font isn't available and a lookalike is used, so
    sizes are scaled per text box to keep line widths close to the PDF.
    """
    if face:
        name, bold, italic = face.run_font
        return name, bold, italic, None, True
    if _anonymous(pdf_name):
        # weight/slanted are measured from the glyphs (see _measure_fonts)
        if mono:
            return "Menlo", weight == "bold", slanted, "cour", False
        if weight == "medium":
            return "Helvetica Neue Medium", False, slanted, "helv", False
        return "Helvetica Neue", weight == "bold", slanted, "hebo" if weight == "bold" else "helv", False

    name = pdf_name.split("+")[-1]
    parts = [p for p in re.split(r"[-,]", name) if p] or ["Helvetica"]
    base, styles = parts[0], [p.lower() for p in parts[1:]]
    bold = bool(set(styles) & BOLD_TOKENS) or bool(flags & 16)
    italic = bool(flags & 2) or any("ital" in s or "oblique" in s for s in styles)
    weight = [w for w in styles if w in WEIGHT_WORDS]
    pretty = re.sub(r"(?<=[a-z])(?=[A-Z])|(?<=[A-Z])(?=[A-Z][a-z])", " ", base)
    pretty = re.sub(r"\s(MT|PS)$", "", pretty)
    low = _norm(base)

    if _is_installed(base):
        family = " ".join([pretty] + [w.capitalize() for w in weight])
        return family, bold, italic, None, True
    if mono or any(k in low for k in ("mono", "courier", "consolas", "menlo", "typewriter", "code")):
        return "Menlo", bold, italic, "cour", False
    if ("serif" in low and "sans" not in low) or any(
            k in low for k in ("times", "georgia", "garamond", "palatino", "minion", "cambria", "bookman")):
        return "Times New Roman", bold, italic, "tibo" if bold else "tiro", False
    return "Helvetica Neue", bold, italic, "hebo" if bold else "helv", False


def _span_font(span):
    return _font_info(span["font"], span["flags"], span["mono"], span["weight"], span["slanted"], span.get("face"))


def _scale_for(lines):
    """Factor so a substituted font occupies about the same width as the PDF's."""
    ratios = []
    for line in lines:
        want = line["bbox"][2] - line["bbox"][0]
        have = 0
        for s in line["spans"]:
            _, _, _, measure, exact = _span_font(s)
            if exact or measure is None:
                return 1.0
            have += pymupdf.Font(measure).text_length(s["text"], s["size"])
            have += s["tracking"] * max(len(s["text"]) - 1, 0) + s["word_extra"] * s["text"].count(" ")
        if have > 0 and want > 0:
            ratios.append(want / have)
    if not ratios:
        return 1.0
    return min(max(statistics.median(ratios), 0.8), 1.15)


# --------------------------------------------------------------------------- #
# Text extraction: rebuild words and letter-spacing from glyph positions
# --------------------------------------------------------------------------- #
def _rebuild(chars, size, direction):
    """Join glyphs into text, inserting spaces at word gaps.

    Returns (text, tracking, word_extra, mono). Tracking is the extra space
    between letters (CSS letter-spacing); word_extra is the extra width each
    real space gets (justified text); mono means every glyph has the same advance.
    """
    dx, dy = direction
    pos = [c["origin"][0] * dx + c["origin"][1] * dy for c in chars]
    adv = [(c["bbox"][2] - c["bbox"][0]) * abs(dx) + (c["bbox"][3] - c["bbox"][1]) * abs(dy) for c in chars]
    gaps = [pos[i + 1] - (pos[i] + adv[i]) for i in range(len(chars) - 1)]
    glyph_pairs = [i for i in range(len(gaps)) if chars[i]["c"] != " " and chars[i + 1]["c"] != " "]

    tracking = statistics.median(gaps[i] for i in glyph_pairs) if len(glyph_pairs) >= 2 else 0
    word_gap = [i for i in glyph_pairs if gaps[i] - tracking > 0.22 * size]
    letter = [gaps[i] for i in glyph_pairs if i not in word_gap]
    if letter:
        tracking = statistics.median(letter)
    # Glyph boxes can be wider than the advance (italic overhang, Type3 font
    # boxes), which reads as negative spacing, so only clear positive
    # letter-spacing is kept; width fitting absorbs the rest.
    tracking = min(tracking, size) if tracking > 0.05 * size else 0

    out = [chars[0]["c"]]
    for i in range(len(gaps)):
        if i in word_gap:
            out.append(" ")
        out.append(chars[i + 1]["c"])

    space_gaps = [gaps[i] - tracking for i in range(len(gaps)) if chars[i]["c"] == " "]
    word_extra = statistics.median(space_gaps) if space_gaps else 0
    word_extra = word_extra if 0.01 * size < word_extra < 2 * size else 0

    widths = [a for a, c in zip(adv, chars) if c["c"].strip()]
    mono = len(widths) >= 4 and max(widths) - min(widths) < 0.03 * size
    return "".join(out), tracking, word_extra, mono


def _merge_spans(spans, hidden):
    """Join consecutive same-style spans (Chrome emits one per word) and drop
    invisible ones, so words and letter-spacing are judged over whole runs."""
    merged = []
    for s in spans:
        if not s["chars"]:
            continue
        sb = pymupdf.Rect(s["bbox"])
        if any(h.contains(sb.tl + (sb.br - sb.tl) * 0.5) for h in hidden):
            continue
        key = (s["font"], round(s["size"], 2), s["color"], s.get("alpha", 255))
        punct = not any(ch["c"].isalnum() for ch in s["chars"])
        if merged and (merged[-1]["key"] == key or (punct and merged[-1]["key"][1:] == key[1:])):
            m = merged[-1]
            m["chars"] = m["chars"] + s["chars"]
            m["bbox"] = tuple(pymupdf.Rect(m["bbox"]) | sb)
        else:
            merged.append(dict(s, key=key, chars=list(s["chars"])))
    return merged


def _split_segments(spans, direction):
    """Split a PDF line where it jumps far sideways or changes baseline: those
    are separate items (a figure and its caption) that only happen to align."""
    if direction != (1.0, 0.0):
        return [spans] if spans else []
    segs = []
    for sp in spans:
        if segs:
            prev = segs[-1][-1]
            size = max(prev["size"], sp["size"])
            far = sp["bbox"][0] - prev["bbox"][2] > 1.2 * size
            shifted = abs(sp["origin"][1] - prev["origin"][1]) > 0.25 * size
            if not (far or shifted):
                segs[-1].append(sp)
                continue
        segs.append([sp])
    return segs


def _baseline(line):
    return max(line["spans"], key=lambda s: s["size"])["origin"][1]


def _extract_blocks(page):
    """Text blocks -> lines -> spans, with spaces, tracking and visibility fixed up."""
    hidden = [pymupdf.Rect(t["bbox"]) for t in page.get_texttrace()
              if t["type"] == 3 or t["opacity"] == 0]
    raw = page.get_text("rawdict", flags=pymupdf.TEXTFLAGS_RAWDICT | pymupdf.TEXT_INHIBIT_SPACES)
    blocks = []
    for b in raw["blocks"]:
        if b["type"] != 0:
            continue
        lines = []
        for l in b["lines"]:
            spans = []
            prev_end = None
            dx, dy = l["dir"]
            for s in _merge_spans(l["spans"], hidden):
                chars = s["chars"]
                text, tracking, word_extra, mono = _rebuild(chars, s["size"], l["dir"])
                start = chars[0]["origin"][0] * dx + chars[0]["origin"][1] * dy
                if prev_end is not None and spans and not text.startswith(" ") \
                        and not spans[-1]["text"].endswith(" ") and start - prev_end > 0.22 * s["size"]:
                    text = " " + text
                last = chars[-1]
                prev_end = last["origin"][0] * dx + last["origin"][1] * dy + \
                    (last["bbox"][2] - last["bbox"][0]) * abs(dx) + (last["bbox"][3] - last["bbox"][1]) * abs(dy)
                spans.append(dict(text=text, size=s["size"], font=s["font"], flags=s["flags"],
                                  color=s["color"], alpha=s.get("alpha", 255), tracking=tracking, word_extra=word_extra,
                                  mono=mono, weight=None, slanted=False, origin=s["origin"], bbox=s["bbox"], chars=chars))
            for seg in _split_segments(spans, l["dir"]):
                if "".join(sp["text"] for sp in seg).strip():
                    bbox = pymupdf.Rect(seg[0]["bbox"])
                    for sp in seg[1:]:
                        bbox |= pymupdf.Rect(sp["bbox"])
                    lines.append(dict(bbox=tuple(bbox), dir=l["dir"], spans=seg))
        if lines:
            blocks.append(lines)
    return blocks


def _spans(pages_blocks):
    for blocks in pages_blocks:
        for lines in blocks:
            for l in lines:
                for s in l["spans"]:
                    yield l, s


def _unify_mono(pages_blocks):
    """Monospace is a property of the font: decide it by majority per font."""
    votes = {}
    for _, s in _spans(pages_blocks):
        if sum(1 for c in s["chars"] if c["c"].strip()) >= 4:
            votes.setdefault(s["font"], []).append(s["mono"])
    mono = {f: sum(v) > len(v) / 2 for f, v in votes.items()}
    for _, s in _spans(pages_blocks):
        s["mono"] = mono.get(s["font"], s["mono"])


SLANTS = np.linspace(-0.4, 0.4, 17)


def _slant_sharpness(a):
    """Deslant test: for each candidate slant, undo it and measure how sharp
    the column profile gets. Upright text peaks at 0, italics near +0.2."""
    h, w = a.shape
    rows = np.arange(h)
    pad = int(np.ceil(0.4 * h)) + 1
    out = np.empty(len(SLANTS))
    for k, t in enumerate(SLANTS):
        offs = np.round(-t * (h / 2 - rows)).astype(int) + pad
        proj = np.zeros(w + 2 * pad)
        for r in range(h):
            proj[offs[r]:offs[r] + w] += a[r]
        out[k] = (proj ** 2).sum()
    return out / out.max()


def _measure_fonts(page, blocks, stats):
    """For each nameless font, accumulate its ink versus regular Helvetica
    drawn with the same glyphs at the same positions, and its slant profile."""
    targets = [s for lines in blocks for l in lines for s in l["spans"]
               if _anonymous(s["font"]) and not s["mono"] and l["dir"] == (1.0, 0.0) and s["text"].strip()]
    if not targets:
        return
    zoom = 3
    src = _page_copy(page)
    sp = src[0]
    sp.add_redact_annot(sp.rect, fill=False)
    sp.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_REMOVE,
                        graphics=pymupdf.PDF_REDACT_LINE_ART_REMOVE_IF_TOUCHED,
                        text=pymupdf.PDF_REDACT_TEXT_NONE)
    ref = pymupdf.open()
    rp = ref.new_page(width=page.rect.width, height=page.rect.height)
    for s in targets:
        for c in s["chars"]:
            if c["c"].strip():
                rp.insert_text(c["origin"], c["c"], fontname="helv", fontsize=s["size"])

    def alpha(p):
        pix = p.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=True)
        return np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.width, pix.n)[:, :, -1].astype(np.float64)

    ink, base = alpha(sp), alpha(rp)
    for s in targets:
        x0, y0, x1, y1 = (int(v * zoom) for v in s["bbox"])
        a, b = ink[y0:y1, x0:x1], base[y0:y1, x0:x1]
        n = sum(1 for c in s["chars"] if c["c"].strip())
        st = stats.setdefault(s["font"], dict(ink=0.0, ref=0.0, n=0, sharp=np.zeros(len(SLANTS)), samples=0))
        st["ink"] += a.sum()
        st["ref"] += b.sum()
        st["n"] += n
        if st["samples"] < 20 and a.shape[0] >= 8 and a.sum() > 0:
            st["sharp"] += _slant_sharpness(a)
            st["samples"] += 1
    src.close()
    ref.close()


def _advance_samples(span):
    """(codepoint, size, advance, run id) for each glyph followed by another."""
    ch = span["chars"]
    out = []
    for a, b in zip(ch, ch[1:]):
        if a["c"] != " " and b["c"] != " ":
            adv = b["origin"][0] - a["origin"][0]
            if 0 < adv < 2 * span["size"]:
                out.append((ord(a["c"]), span["size"], adv, id(span)))
    return out


SHAPE_ZOOM = 4


def _ink(page, clip):
    pix = page.get_pixmap(matrix=pymupdf.Matrix(SHAPE_ZOOM, SHAPE_ZOOM), clip=clip, alpha=True)
    return np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.width, pix.n)[:, :, -1].astype(float)


def _pdf_ink(page, span):
    """The span's glyphs as the PDF draws them (text only), plus the clip."""
    clip = pymupdf.Rect(span["bbox"]) + (-1, -1, 1, 1)
    src = _page_copy(page)
    sp = src[0]
    sp.add_redact_annot(sp.rect, fill=False)
    sp.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_REMOVE,
                        graphics=pymupdf.PDF_REDACT_LINE_ART_REMOVE_IF_TOUCHED,
                        text=pymupdf.PDF_REDACT_TEXT_NONE)
    ink = _ink(sp, clip)
    src.close()
    return ink, clip


def _shape_score(page, span, target, clip, face):
    """Overlap between the PDF's glyphs and the same glyphs drawn with `face`
    at the same positions (1 = identical)."""
    if face.index or face.variable:
        return 0.0  # can't draw a .ttc member or a variable instance here
    ref = pymupdf.open()
    rp = ref.new_page(width=page.rect.width, height=page.rect.height)
    try:
        for c in span["chars"]:
            if c["c"].strip():
                rp.insert_text(c["origin"], c["c"], fontsize=span["size"], fontname="F0", fontfile=face.path)
        ink = _ink(rp, clip)
    except Exception:
        return 0.0
    finally:
        ref.close()
    if ink.shape != target.shape or max(ink.sum(), target.sum()) == 0:
        return 0.0
    return float(np.minimum(ink, target).sum() / np.maximum(ink, target).sum())


def _pick_by_shape(page, span, tied, families_seen):
    """Among faces whose widths fit equally well, the one whose glyphs look
    most like the PDF's. A family already chosen for this document wins ties
    again, so the comparison runs once per family rather than per page."""
    known = [m for m in tied if m.face.family in families_seen]
    if known:
        tied = known
    target, clip = _pdf_ink(page, span)
    by_family = {}
    for m in tied:
        by_family.setdefault(m.face.family, []).append(m)
    if len(by_family) > 1:
        # choose the family from its common weights first, then the weight
        common = ("Regular", "Medium", "Bold", "Italic", "Bold Italic")
        best_of = {f: max((m for m in ms if m.face.style in common), default=ms[0],
                          key=lambda m: _shape_score(page, span, target, clip, m.face))
                   for f, ms in list(by_family.items())[:12]}
        family = max(best_of, key=lambda f: _shape_score(page, span, target, clip, best_of[f].face))
        tied = by_family[family]
    return max(tied[:24], key=lambda m: _shape_score(page, span, target, clip, m.face))


def _identify_fonts(doc, pages_blocks):
    """Name each nameless font by matching its glyph widths against installed
    faces; matched spans get the exact face and exact letter-spacing. When
    several families fit (monospace fonts all share one width), the glyph
    shapes decide."""
    samples, slant, longest = {}, {}, {}
    for page_no, blocks in enumerate(pages_blocks):
        for l, s in _spans([blocks]):
            if _anonymous(s["font"]) and l["dir"] == (1.0, 0.0):
                samples.setdefault(s["font"], []).extend(_advance_samples(s))
                slant.setdefault(s["font"], s["slanted"] if s["weight"] else None)
                if len(s["chars"]) > len(longest.get(s["font"], (0, {"chars": []}))[1]["chars"]):
                    longest[s["font"]] = (page_no, s)
    matches, families_seen = {}, set()
    for f in sorted(samples, key=lambda f: -len(samples[f])):
        tied = fonts.identify(samples[f][:300], slant.get(f))
        if not tied:
            continue
        if len({m.face.family for m in tied}) > 1 or len(tied) > 1 and tied[0].face.mono:
            page_no, span = longest[f]
            best = _pick_by_shape(doc[page_no], span, tied, families_seen)
        else:
            best = tied[0]
        families_seen.add(best.face.family)
        matches[f] = best
    for l, s in _spans(pages_blocks):
        m = matches.get(s["font"])
        if not m:
            continue
        s["face"] = m.face
        # letter-spacing from this span's own glyphs against the real widths
        res = [adv - m.face.adv[cp] * size for cp, size, adv, _ in _advance_samples(s) if cp in m.face.adv]
        t = statistics.median(res) if len(res) >= 2 else m.tracking * s["size"]
        s["tracking"] = t if abs(t) > 0.004 * s["size"] else 0
    return matches


def _assign_styles(pages_blocks, stats):
    """Weight relative to the document's body font (the one with the most
    glyphs), which cancels out how heavy the family is compared to Helvetica:
    within ~15% is regular, up to ~45% heavier is medium, beyond that bold."""
    ratios = {f: st["ink"] / st["ref"] for f, st in stats.items() if st["ref"] > 0 and st["n"] >= 3}
    if not ratios:
        return
    body = ratios[max(ratios, key=lambda f: stats[f]["n"])]
    weight, slanted = {}, {}
    for f, r in ratios.items():
        rel = r / body
        weight[f] = "bold" if rel > 1.45 or r > 1.75 else "medium" if rel > 1.15 else "regular"
        st = stats[f]
        slanted[f] = bool(st["samples"] > 0 and SLANTS[int(np.argmax(st["sharp"]))] >= 0.1)
    for _, s in _spans(pages_blocks):
        if s["font"] in weight:
            s["weight"] = weight[s["font"]]
            s["slanted"] = slanted[s["font"]]


# --------------------------------------------------------------------------- #
# Art layer (default) and separate pictures (optional)
# --------------------------------------------------------------------------- #
def _add_art(slide, page, fit):
    """The page with its text removed, as one high-resolution picture."""
    copy = _page_copy(page)
    cp = copy[0]
    cp.add_redact_annot(cp.rect, fill=False)
    cp.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,
                        graphics=pymupdf.PDF_REDACT_LINE_ART_NONE,
                        text=pymupdf.PDF_REDACT_TEXT_REMOVE)
    zoom = min(ART_WIDTH_PX / page.rect.width, 8)
    pix = cp.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
    data = pix.tobytes("jpg", jpg_quality=92) if page.get_images() else pix.tobytes("png")
    copy.close()
    slide.shapes.add_picture(BytesIO(data), fit.x(0), fit.y(0), fit.d(page.rect.width), fit.d(page.rect.height))


def _image_data(doc, xref):
    """Original bytes when the image is plain RGB/gray JPEG/PNG, else re-encode."""
    try:
        info = doc.extract_image(xref)
    except Exception:
        info = None
    if info and not info.get("smask") and info["ext"] in ("jpeg", "jpg", "png") \
            and info.get("colorspace") in (1, 3):
        return info["image"]

    pix = pymupdf.Pixmap(doc, xref)
    if pix.colorspace and pix.colorspace.n >= 4:  # CMYK -> RGB
        pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
    smask = doc.xref_get_key(xref, "SMask")
    if smask[0] == "xref":
        mask = pymupdf.Pixmap(doc, int(smask[1].split()[0]))
        if (mask.width, mask.height) == (pix.width, pix.height):
            if pix.alpha:
                pix = pymupdf.Pixmap(pix, 0)
            pix = pymupdf.Pixmap(pix, mask)
    return pix.tobytes("png")


def _add_images(slide, page, fit):
    doc = page.parent
    for info in page.get_image_info(xrefs=True):
        xref = info.get("xref")
        if not xref:
            continue
        try:
            data = _image_data(doc, xref)
        except Exception as e:
            print(f"  skipped image {xref}: {e}", file=sys.stderr)
            continue
        r = pymupdf.Rect(info["bbox"]).normalize()
        slide.shapes.add_picture(BytesIO(data), fit.x(r.x0), fit.y(r.y0), fit.d(r.width), fit.d(r.height))


def _add_line_art(slide, page, fit):
    """Everything that isn't text or an image, as one transparent picture."""
    drawings = page.get_drawings()
    if not drawings:
        return
    copy = _page_copy(page)
    cp = copy[0]

    # Full-page fills (the page "paper") would make the overlay opaque and hide
    # the pictures below it. Use the topmost one as the slide background and
    # strip them from the overlay; they all touch the corner.
    page_area = page.rect.width * page.rect.height
    full = [d["fill"] for d in drawings if d.get("fill") and d["rect"].get_area() > 0.95 * page_area]
    if full:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(*(round(c * 255) for c in full[-1]))
        cp.add_redact_annot(pymupdf.Rect(0, 0, 1, 1), fill=False)
        cp.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,
                            graphics=pymupdf.PDF_REDACT_LINE_ART_REMOVE_IF_TOUCHED,
                            text=pymupdf.PDF_REDACT_TEXT_NONE)
        if not cp.get_drawings():
            copy.close()
            return

    cp.add_redact_annot(cp.rect, fill=False)
    cp.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_REMOVE,
                        graphics=pymupdf.PDF_REDACT_LINE_ART_NONE)
    zoom = min(ART_WIDTH_PX / page.rect.width, 8)
    data = cp.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=True).tobytes("png")
    copy.close()
    slide.shapes.add_picture(BytesIO(data), fit.x(0), fit.y(0), fit.d(page.rect.width), fit.d(page.rect.height))


# --------------------------------------------------------------------------- #
# Text boxes
# --------------------------------------------------------------------------- #
def _is_rotated(line):
    dx, dy = line["dir"]
    return abs(dy) > 0.02 or dx < 0


def _alignment(lines):
    if len(lines) < 2:
        return PP_ALIGN.LEFT
    x0 = [l["bbox"][0] for l in lines]
    x1 = [l["bbox"][2] for l in lines]
    mid = [(a + b) / 2 for a, b in zip(x0, x1)]
    if max(x0) - min(x0) < 2:
        return PP_ALIGN.LEFT
    if max(x1) - min(x1) < 2:
        return PP_ALIGN.RIGHT
    if max(mid) - min(mid) < 2:
        return PP_ALIGN.CENTER
    return PP_ALIGN.LEFT


def _split_lines(lines):
    """PDF text blocks often hold unrelated text (a row of labels, a callout
    label). Start a new group when a line sits beside, or far off to the side
    of, the group so far; one-column paragraphs stay together. Rotated lines
    each get their own group.
    """
    flat = [l for l in lines if not _is_rotated(l)]
    groups = [[l] for l in lines if _is_rotated(l)]
    if _alignment(flat) != PP_ALIGN.LEFT:
        return groups + ([flat] if flat else [])
    cur = None
    for line in flat:
        size = max(s["size"] for s in line["spans"])
        if cur:
            prev_size = max(s["size"] for s in cur[-1]["spans"])
            same_column = abs(line["bbox"][0] - cur[0]["bbox"][0]) <= 3 * size
            below = _baseline(line) > _baseline(cur[-1]) + 0.3 * size
            close = _baseline(line) - _baseline(cur[-1]) < 2.5 * size
            same_size = abs(size - prev_size) <= 0.15 * max(size, prev_size)
            if same_column and below and close and same_size:
                cur.append(line)
                continue
        cur = [line]
        groups.append(cur)
    return groups


def _set_color(font, color, alpha):
    font.color.rgb = RGBColor((color >> 16) & 255, (color >> 8) & 255, color & 255)
    if alpha < 255:
        clr = font.color._color._xClr  # <a:srgbClr>
        el = clr.makeelement(qn("a:alpha"), {"val": str(round(alpha / 255 * 100000))})
        clr.append(el)


# Keynote with exact line spacing L puts a line's baseline 0.8*L below the top
# of that line, and each line is L tall, for every font. Measured, see README.
BASELINE_IN_LINE = 0.8


def _line_heights(lines):
    """Exact line heights (PDF pt) that put every baseline where the PDF has it."""
    baselines = [_baseline(l) for l in lines]
    first = max(s["size"] for s in lines[0]["spans"]) * 1.2
    gaps = [b - a for a, b in zip(baselines, baselines[1:])]
    if gaps:
        first = statistics.median(gaps)
    heights = [first]
    for gap, line in zip(gaps, lines[1:]):
        size = max(s["size"] for s in line["spans"])
        # baseline_i = baseline_{i-1} + (1 - 0.8) * L_{i-1} + 0.8 * L_i
        h = (gap - (1 - BASELINE_IN_LINE) * heights[-1]) / BASELINE_IN_LINE
        heights.append(max(h, 0.5 * size))
    return baselines[0], heights


def _add_text_box(slide, lines, fit, stats=None):
    x0 = min(l["bbox"][0] for l in lines)
    y0 = min(l["bbox"][1] for l in lines)
    x1 = max(l["bbox"][2] for l in lines)
    y1 = max(l["bbox"][3] for l in lines)
    scale = _scale_for(lines)
    align = _alignment(lines)
    if stats is not None:
        stats["boxes"] += 1
        if abs(scale - 1) > 0.005:
            stats["scales"].append(scale)
        stats["rotated"] += _is_rotated(lines[0])
        stats["letter_spaced"] += sum(1 for l in lines for sp in l["spans"] if sp["tracking"])
        stats["justified_lines"] += sum(1 for l in lines if any(sp["word_extra"] for sp in l["spans"]))
    rotation = 0
    width = (x1 - x0) * 1.02 + 4

    if _is_rotated(lines[0]):
        dx, dy = lines[0]["dir"]
        rotation = math.degrees(math.atan2(dy, dx)) % 360
        c, s = abs(dx), abs(dy)
        size = lines[0]["spans"][0]["size"]
        heights = [size * 1.2]
        length = max((x1 - x0) * c + (y1 - y0) * s - 2 * heights[0] * s * c, heights[0])
        cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
        width, height = length + 4, heights[0]
        left, top = cx - width / 2, cy - height / 2
    else:
        baseline, heights = _line_heights(lines)
        top = baseline - BASELINE_IN_LINE * heights[0]
        height = sum(heights)
        left = {PP_ALIGN.RIGHT: x1 - width, PP_ALIGN.CENTER: (x0 + x1) / 2 - width / 2}.get(align, x0)

    box = slide.shapes.add_textbox(fit.x(left), fit.y(top), fit.d(width), fit.d(height))
    box.rotation = rotation
    tf = box.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = Pt(heights[min(i, len(heights) - 1)] * fit.s)
        para.space_before = para.space_after = Pt(0)
        for span in line["spans"]:
            if not span["text"]:
                continue
            _add_runs(para, span, scale, fit)


def _add_runs(para, span, scale, fit):
    """One run per span; justified lines get their spaces as separate runs
    widened by the PDF's extra word spacing, since Keynote won't justify
    lines that don't wrap."""
    family, bold, italic, _, exact = _span_font(span)
    k = (1 if exact else scale) * fit.s
    pieces = re.split(r"( +)", span["text"]) if span["word_extra"] else [span["text"]]
    for piece in pieces:
        if not piece:
            continue
        spacing = span["tracking"]
        if piece.startswith(" ") and span["word_extra"]:
            spacing += span["word_extra"]
        run = para.add_run()
        run.text = piece
        font = run.font
        font.size = Pt(max(span["size"] * k, 1))
        font.name = family
        font.bold = bold
        font.italic = italic
        if spacing:
            run._r.get_or_add_rPr().set("spc", str(round(spacing * fit.s * 100)))
        _set_color(font, span["color"], span["alpha"])


def _add_text(slide, blocks, fit, stats=None):
    for lines in blocks:
        for group in _split_lines(lines):
            _add_text_box(slide, group, fit, stats)


# --------------------------------------------------------------------------- #
# Build + Keynote
# --------------------------------------------------------------------------- #
def open_pdf(pdf_path):
    doc = pymupdf.open(pdf_path)
    if doc.needs_pass:
        raise ValueError("PDF is password protected")
    if not len(doc):
        raise ValueError("PDF has no pages")
    return doc


def new_presentation(width_pt, height_pt):
    prs = Presentation()
    prs.slide_width, prs.slide_height = _emu(width_pt), _emu(height_pt)
    return prs


def analyze_text(doc, progress=None, report=None, span=(0.0, 0.3)):
    """Text of every page, rebuilt and with fonts resolved: one block list per page.
    `span` is the share of the overall progress bar this step covers."""
    lo, hi = span

    def say(msg, frac):
        if progress:
            progress(msg, lo + (hi - lo) * frac)

    pages_blocks, stats = [], {}
    n = len(doc)
    for i, page in enumerate(doc):
        say(f"Reading text on page {i + 1} of {n}", 0.4 * i / n)
        pages_blocks.append(_extract_blocks(page))
    _unify_mono(pages_blocks)
    for i, (page, blocks) in enumerate(zip(doc, pages_blocks)):
        say(f"Measuring fonts on page {i + 1} of {n}", 0.4 + 0.4 * i / n)
        _measure_fonts(page, blocks, stats)
    _assign_styles(pages_blocks, stats)
    say("Recognising fonts from their letter widths", 0.85)
    _identify_fonts(doc, pages_blocks)
    reporting.record_fonts(report, pages_blocks, _span_font, _anonymous)
    return pages_blocks


def add_text(slide, blocks, page_rect, slide_w, slide_h, report=None):
    _add_text(slide, blocks, _Fit(page_rect, slide_w, slide_h), reporting.text_stats(report))


def pdf_to_pptx(pdf_path, pptx_path, separate_images=False, progress=None, report=None):
    doc = open_pdf(pdf_path)
    first = doc[0].rect
    prs = new_presentation(first.width, first.height)
    blank = prs.slide_layouts[6]
    pages_blocks = analyze_text(doc, progress, report)

    for i, page in enumerate(doc):
        if progress:
            progress(f"Building slide {i + 1} of {len(doc)}", 0.3 + 0.4 * i / len(doc))
        fit = _Fit(page.rect, first.width, first.height)
        slide = prs.slides.add_slide(blank)
        if separate_images:
            _add_images(slide, page, fit)
            _add_line_art(slide, page, fit)
        else:
            _add_art(slide, page, fit)
        _add_text(slide, pages_blocks[i], fit, reporting.text_stats(report))
    reporting.record_objects(report, prs)
    prs.save(pptx_path)


def _wait_for_keynote(timeout=30):
    subprocess.run(["open", "-a", "Keynote"], check=True)  # `launch` via AppleScript can fail
    end = time.time() + timeout
    while time.time() < end:
        if subprocess.run(["pgrep", "-x", "Keynote"], capture_output=True).returncode == 0:
            time.sleep(2)
            return
        time.sleep(0.5)
    raise RuntimeError("Keynote did not start")


def pptx_to_keynote(pptx_path, key_path, png_dir=None):
    """Open the .pptx in Keynote, save it as .key and optionally export slide PNGs."""
    script = """
    on run argv
        set inFile to POSIX file (item 1 of argv)
        set outFile to POSIX file (item 2 of argv)
        tell application "Keynote"
            open inFile
            -- `open` can return an import-warnings object instead of the
            -- document, so look the document up by name once it appears
            set d to missing value
            repeat 150 times
                try
                    set d to first document whose name is (item 4 of argv)
                    exit repeat
                end try
                delay 0.2
            end repeat
            if d is missing value then error "Keynote did not open the presentation"
            save d in outFile
            if (item 3 of argv) is not "" then
                export d to POSIX file (item 3 of argv) as slide images
            end if
            close d saving no
        end tell
    end run
    """
    # Keynote resolves relative paths against its own working directory (/),
    # and crashes when it can't write there, so always pass absolute paths.
    args = ["osascript", "-e", script, str(Path(pptx_path).resolve()), str(Path(key_path).resolve()),
            str(Path(png_dir).resolve()) if png_dir else "", Path(pptx_path).stem]
    for attempt in range(3):
        _wait_for_keynote()
        r = subprocess.run(args, capture_output=True, text=True, timeout=900)
        if not r.returncode:
            return
        # -600/-609: Keynote quit or crashed. -1700/-1708: its scripting got
        # into a bad state (seen after crashes). A clean relaunch fixes both.
        if attempt < 2 and re.search(r"\((-600|-609|-1700|-1708)\)", r.stderr):
            subprocess.run(["osascript", "-e", 'tell application "Keynote" to quit saving no'],
                           capture_output=True, timeout=30)
            time.sleep(3)
            subprocess.run(["pkill", "-x", "Keynote"], capture_output=True)
            time.sleep(1)
            # start fresh: restoring windows from a crashed session can wedge it again
            subprocess.run(["open", "-F", "-g", "-a", "Keynote"], capture_output=True)
            continue
        raise RuntimeError(f"Keynote conversion failed: {r.stderr.strip()}")


# --------------------------------------------------------------------------- #
# Accuracy check
# --------------------------------------------------------------------------- #
@dataclass
class SlideReport:
    page: int
    match: float       # 0-100, higher is better
    off_pixels: float  # % of pixels that differ noticeably


def compare(pdf_path, png_dir, out_dir, thumb_width=720):
    """Compare each exported Keynote slide with the PDF page; save thumbnails."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    pngs = sorted(Path(png_dir).glob("*.png"))
    doc = pymupdf.open(pdf_path)
    reports = []
    for i, page in enumerate(doc):
        if i >= len(pngs):
            break
        key_img = Image.open(pngs[i]).convert("RGB")
        w, h = key_img.size
        # Render the page the way the slide shows it: fit and centered.
        first = doc[0].rect
        fit = _Fit(page.rect, first.width, first.height)
        k = w / first.width
        pix = page.get_pixmap(matrix=pymupdf.Matrix(fit.s * k, fit.s * k), alpha=False)
        pdf_img = Image.new("RGB", (w, h), "white")
        pdf_img.paste(Image.frombytes("RGB", (pix.width, pix.height), pix.samples),
                      (round(fit.ox * k), round(fit.oy * k)))
        diff = np.abs(np.asarray(pdf_img, dtype=np.int16) - np.asarray(key_img, dtype=np.int16))
        match = 100 * (1 - diff.mean() / 255)
        off = 100 * (diff.max(axis=2) > 48).mean()
        reports.append(SlideReport(i + 1, round(match, 2), round(off, 2)))

        heat = Image.fromarray(np.clip(diff.max(axis=2) * 3, 0, 255).astype(np.uint8)).convert("RGB")
        th = round(thumb_width * h / w)
        for name, im in (("pdf", pdf_img), ("key", key_img), ("diff", heat)):
            im.resize((thumb_width, th)).save(out_dir / f"{i + 1}_{name}.jpg", quality=85)
    return reports


def missing_fonts(pdf_path):
    """Font families used by the PDF that aren't installed (they get substituted)."""
    missing = set()
    doc = pymupdf.open(pdf_path)
    for page in doc:
        for f in page.get_fonts():
            if _anonymous(f[3]):
                continue  # identified from glyph widths during conversion
            _, _, _, _, exact = _font_info(f[3], 0)
            if not exact:
                base = re.split(r"[-,]", f[3].split("+")[-1])[0]
                missing.add(re.sub(r"(?<=[a-z])(?=[A-Z])|(?<=[A-Z])(?=[A-Z][a-z])", " ", base))
    return sorted(missing)


# --------------------------------------------------------------------------- #
# Public entry point
# --------------------------------------------------------------------------- #
def _finish(pptx, key_path, pptx_out, keynote, verify, original_pdf, progress, preview_dir, tmp):
    """Shared tail of both inputs: keep the .pptx, make the .key, compare."""
    if pptx_out:
        Path(pptx_out).parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(pptx, pptx_out)
    if not keynote:
        return []
    if progress:
        progress("Importing into Keynote", 0.8)
    png_dir = tmp / "slides" if verify else None
    pptx_to_keynote(pptx, key_path, png_dir)
    if not verify:
        return []
    if progress:
        progress("Comparing slides with the original", 0.95)
    return compare(original_pdf, png_dir, preview_dir or tmp / "preview")


def convert(pdf_path, key_path, separate_images=False, verify=True, progress=None, preview_dir=None,
            report=None, pptx_out=None, keynote=True):
    """PDF -> .key and/or .pptx. Returns a list of SlideReport (empty if verify
    or keynote is off; the check needs Keynote to render the slides).

    pptx_out: also save the PowerPoint file there. keynote=False skips Keynote
    entirely (key_path is then ignored). Pass a dict from report.new() as
    `report` to receive the status report."""
    started = time.time()
    pdf_path = Path(pdf_path).resolve()
    key_path = Path(key_path).resolve() if key_path else None
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        pptx = tmp / "deck.pptx"
        pdf_to_pptx(pdf_path, pptx, separate_images, progress, report)
        slides = _finish(pptx, key_path, pptx_out, keynote, verify, pdf_path, progress, preview_dir, tmp)
        reporting.record_accuracy(report, slides, time.time() - started)
        return slides

"""Tests for the PDF -> .pptx stage (no Keynote needed)."""
import pymupdf
import pytest
from pptx import Presentation

from pdf2key import pdf_to_pptx
from pdf2key.converter import _font_info, _rebuild


@pytest.fixture
def sample_pdf(tmp_path):
    doc = pymupdf.open()
    page = doc.new_page(width=800, height=450)
    page.insert_text((50, 80), "Title", fontname="tiro", fontsize=40)
    page.insert_text((50, 140), "Line one", fontsize=18, color=(1, 0, 0))
    page.insert_text((50, 162), "Line two", fontsize=18, color=(1, 0, 0))
    page.insert_text((400, 300), "Rotated", fontsize=20, rotate=90)
    page.draw_line((50, 200), (700, 200), color=(0, 0, 1), width=2)
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 40), False)
    pix.set_rect(pix.irect, (0, 200, 0))
    page.insert_image(pymupdf.Rect(600, 300, 700, 400), pixmap=pix)
    portrait = doc.new_page(width=450, height=800)
    portrait.insert_text((50, 100), "Portrait", fontsize=30)
    path = tmp_path / "sample.pdf"
    doc.save(path)
    return path


def _texts(slide):
    return [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]


def test_builds_one_slide_per_page(sample_pdf, tmp_path):
    out = tmp_path / "out.pptx"
    pdf_to_pptx(sample_pdf, out)
    prs = Presentation(out)
    assert len(prs.slides) == 2
    assert prs.slide_width == 800 * 12700 and prs.slide_height == 450 * 12700


def test_text_is_editable_and_grouped(sample_pdf, tmp_path):
    out = tmp_path / "out.pptx"
    pdf_to_pptx(sample_pdf, out)
    texts = _texts(Presentation(out).slides[0])
    assert "Title" in texts
    assert "Line one\nLine two" in texts  # consecutive lines stay one text box
    assert "Rotated" in texts


def test_rotated_text_keeps_rotation(sample_pdf, tmp_path):
    out = tmp_path / "out.pptx"
    pdf_to_pptx(sample_pdf, out)
    box = next(s for s in Presentation(out).slides[0].shapes if s.has_text_frame and s.text_frame.text == "Rotated")
    assert round(box.rotation) % 90 == 0 and round(box.rotation) != 0


def test_art_layer_is_one_picture_by_default(sample_pdf, tmp_path):
    out = tmp_path / "out.pptx"
    pdf_to_pptx(sample_pdf, out)
    pictures = [s for s in Presentation(out).slides[0].shapes if s.shape_type == 13]
    assert len(pictures) == 1  # everything except text, rendered as the PDF draws it


def test_separate_images_mode(sample_pdf, tmp_path):
    out = tmp_path / "out.pptx"
    pdf_to_pptx(sample_pdf, out, separate_images=True)
    pictures = [s for s in Presentation(out).slides[0].shapes if s.shape_type == 13]
    assert len(pictures) == 2  # the embedded image + the line-art overlay


def test_smaller_page_is_fit_inside_slide(sample_pdf, tmp_path):
    out = tmp_path / "out.pptx"
    pdf_to_pptx(sample_pdf, out)
    prs = Presentation(out)
    for shape in prs.slides[1].shapes:
        assert 0 <= shape.left and shape.left + shape.width <= prs.slide_width + 12700 * 8
        assert 0 <= shape.top and shape.top + shape.height <= prs.slide_height


@pytest.mark.parametrize("name, family_kind", [
    ("ABCDEF+SomeUnknownMono-Bold", "Menlo"),
    ("XYZ+UnknownSerifPro-Regular", "Times New Roman"),
    ("Q+TotallyUnknownSans", "Helvetica Neue"),
])
def test_font_fallbacks(name, family_kind):
    family, *_ , exact = _font_info(name, 0)
    assert not exact and family == family_kind


def test_bold_italic_detection():
    _, bold, italic, *_ = _font_info("ABC+Unknown-BoldItalic", 0)
    assert bold and italic


def _glyphs(text, advance, gap, size=10):
    """Fake rawdict chars: fixed advance, `gap` extra space after each glyph."""
    chars, x = [], 0.0
    for ch in text:
        chars.append(dict(c=ch, origin=(x, 0), bbox=(x, -size, x + advance, 0)))
        x += advance + gap
    return chars


def test_rebuild_restores_word_gaps_and_letter_spacing():
    # "AB CD" drawn without a space glyph, letters spaced 2pt apart, word gap 8pt
    chars = _glyphs("AB", 5, 2) + [dict(c=c["c"], origin=(c["origin"][0] + 20, 0),
                                        bbox=(c["bbox"][0] + 20, -10, c["bbox"][2] + 20, 0))
                                   for c in _glyphs("CD", 5, 2)]
    text, tracking, _, mono = _rebuild(chars, 10, (1.0, 0.0))
    assert text == "AB CD"
    assert tracking == pytest.approx(2)
    assert mono


def test_rebuild_ignores_negative_spacing():
    text, tracking, _, _ = _rebuild(_glyphs("Hello", 5, -1), 10, (1.0, 0.0))
    assert text == "Hello" and tracking == 0


def test_text_baseline_matches_pdf(sample_pdf, tmp_path):
    from pdf2key.converter import BASELINE_IN_LINE
    out = tmp_path / "out.pptx"
    pdf_to_pptx(sample_pdf, out)
    box = next(s for s in Presentation(out).slides[0].shapes if s.has_text_frame and s.text_frame.text == "Title")
    line = box.text_frame.paragraphs[0].line_spacing.pt
    assert (box.top / 12700) + BASELINE_IN_LINE * line == pytest.approx(80, abs=0.5)


def test_rebuild_measures_justified_word_spacing():
    # "ab cd" with a real space glyph followed by 3pt of extra (justified) space
    chars = _glyphs("ab ", 5, 0)
    chars += [dict(c=c["c"], origin=(c["origin"][0] + 18, 0), bbox=(c["bbox"][0] + 18, -10, c["bbox"][2] + 18, 0))
              for c in _glyphs("cd", 5, 0)]
    text, tracking, word_extra, _ = _rebuild(chars, 10, (1.0, 0.0))
    assert text == "ab cd" and tracking == 0 and word_extra == pytest.approx(3)


def _face(family, style, adv, mono=False, italic=False, system=False):
    from pdf2key.fonts import Face
    return Face(family, style, 400, italic, mono, False, system, "", 0, adv)


def test_identify_font_from_glyph_widths():
    from pdf2key.fonts import identify
    narrow = {cp: 0.5 for cp in range(32, 127)}
    wide = {cp: 0.5 + (cp % 7) * 0.05 for cp in range(32, 127)}
    faces = [_face("Narrow", "Regular", narrow), _face("Wide", "Bold", wide)]
    # "Wide" drawn at 20pt with 1pt letter-spacing, in two runs
    text = "Hamburgefonstiv"
    samples = [(ord(c), 20, wide[ord(c)] * 20 + 1, run) for run in (1, 2) for c in text]
    best = identify(samples, faces=faces)[0]
    assert best.face.family == "Wide" and best.face.keynote_name == "Wide Bold"
    assert best.tracking == pytest.approx(1 / 20)


def test_identify_prefers_user_installed_and_italic_hint():
    from pdf2key.fonts import identify
    mono = {cp: 0.6 for cp in range(32, 127)}
    faces = [_face("SystemMono", "Regular", mono, mono=True, system=True),
             _face("UserMono", "Regular", mono, mono=True),
             _face("UserMono", "Italic", mono, mono=True, italic=True)]
    samples = [(ord(c), 10, 6.0, 0) for c in "ABCDEFGH"]
    assert identify(samples, slanted=False, faces=faces)[0].face.keynote_name == "UserMono"
    best = identify(samples, slanted=True, faces=faces)[0].face
    assert best.family == "UserMono" and best.italic


def test_cli_format_pptx_needs_no_keynote(sample_pdf, tmp_path, capsys):
    from pdf2key.cli import main
    main([str(sample_pdf), str(tmp_path / "out.key"), "--format", "pptx"])
    assert (tmp_path / "out.pptx").exists() and not (tmp_path / "out.key").exists()
    assert len(Presentation(tmp_path / "out.pptx").slides) == 2
    assert "Conversion report" in capsys.readouterr().out

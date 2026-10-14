"""Tests for the experimental HTML deck input."""
import pytest
from PIL import Image
from pptx import Presentation

from pdf2key.html import _displayed_rect

EMU_PER_PX = 12700 * 0.75


def test_object_fit_cover_crops_the_overflow():
    # 400x200 image in a 100x100 box, cover, centred: drawn 200x100, 50px off each side
    item = {"kind": "image", "natural": [400, 200], "fit": "cover", "pos": "50% 50%",
            "content": {"x": 0, "y": 0, "w": 100, "h": 100}}
    assert _displayed_rect(item) == pytest.approx((-50, 0, 200, 100))


def test_background_size_contain_and_position():
    item = {"kind": "bgimage", "natural": [200, 100], "size": "contain", "pos": "0% 100%",
            "rect": {"x": 10, "y": 10, "w": 100, "h": 100}}
    assert _displayed_rect(item) == pytest.approx((10, 60, 100, 50))


@pytest.fixture
def chrome():
    pw = pytest.importorskip("playwright.sync_api")
    try:
        with pw.sync_playwright() as p:
            p.chromium.launch(channel="chrome").close()
    except Exception:
        pytest.skip("Google Chrome not available")


def test_html_deck_gets_movable_cropped_pictures_under_overlays(chrome, tmp_path):
    from pdf2key.html import html_to_pptx
    Image.new("RGB", (800, 400), (200, 40, 40)).save(tmp_path / "photo.png")
    (tmp_path / "deck.html").write_text("""<!doctype html><html><body>
      <section style="position:relative;width:1920px;height:1080px;background:#111">
        <img src="photo.png" style="position:absolute;left:100px;top:100px;width:400px;height:400px;object-fit:cover">
        <div style="position:absolute;left:100px;top:100px;width:400px;height:400px;
                    background:linear-gradient(transparent,#000)"></div>
        <h1 style="position:absolute;left:600px;top:100px;color:#fff;font:64px Helvetica">Hello deck</h1>
      </section>
      <section style="position:relative;width:1920px;height:1080px;background:#fff"><p>Second</p></section>
    </body></html>""")
    out = tmp_path / "deck.pptx"
    html_to_pptx(tmp_path / "deck.html", out, tmp_path / "deck.pdf")
    slides = Presentation(out).slides
    assert len(slides) == 2
    shapes = list(slides[0].shapes)
    # the photo and the gradient lying only on it are grouped: photo first, overlay above
    group = next(s for s in shapes if s.shape_type == 6)
    photo, overlay = list(group.shapes)
    assert photo.name == "Photo · photo.png" and overlay.name == "Overlay"
    assert photo.crop_left == pytest.approx(0.25, abs=0.01) and photo.crop_right == pytest.approx(0.25, abs=0.01)
    assert photo.left == pytest.approx(100 * EMU_PER_PX, rel=0.01)
    assert shapes[0].name == "Background"
    assert any(s.has_text_frame and "Hello deck" in s.text_frame.text for s in shapes)


def test_decoration_layers_split_into_separate_pieces():
    from pdf2key.html import SCALE, _pieces
    img = Image.new("RGBA", (400 * SCALE, 200 * SCALE), (0, 0, 0, 0))
    img.paste((255, 0, 0, 255), (10 * SCALE, 10 * SCALE, 60 * SCALE, 60 * SCALE))     # a card
    img.paste((0, 0, 255, 255), (200 * SCALE, 100 * SCALE, 300 * SCALE, 104 * SCALE))  # an arrow shaft
    pieces = _pieces(img)
    assert len(pieces) == 2
    assert sorted(p[0][0] // SCALE for p in pieces) == [10, 200]


def test_find_html_accepts_zip_folder_or_file(tmp_path):
    import zipfile
    from pdf2key.html import _find_html
    deck = tmp_path / "export" / "My Deck"
    (deck / "assets").mkdir(parents=True)
    (deck / "My Deck.html").write_text("<section>" + "x" * 500 + "</section>")
    (deck / "assets" / "helper.html").write_text("<p>small</p>")
    assert _find_html(deck.parent, tmp_path / "w1").name == "My Deck.html"       # folder
    assert _find_html(deck / "My Deck.html", tmp_path / "w2").name == "My Deck.html"  # file
    z = tmp_path / "deck.zip"
    with zipfile.ZipFile(z, "w") as zf:
        for f in deck.rglob("*.html"):
            zf.write(f, f.relative_to(tmp_path))
    (tmp_path / "w3").mkdir()
    assert _find_html(z, tmp_path / "w3").name == "My Deck.html"                  # zip
    with pytest.raises(ValueError):
        _find_html(tmp_path / "deck.txt", tmp_path / "w4")


def test_list_decks_ranks_the_real_deck_first(tmp_path):
    from pdf2key.html import list_decks
    for name, n in [("Deck.html", 30), ("Deck v3.html", 50), ("Deck v3 - backup.html", 55),
                    ("Deck v3-print.html", 50), ("ref/Other.html", 60), ("notes.html", 0)]:
        f = tmp_path / name
        f.parent.mkdir(exist_ok=True)
        f.write_text("<section></section>" * n)
    names = [d for d, _ in list_decks(tmp_path)]
    assert names[:2] == ["Deck v3.html", "Deck.html"]
    assert "notes.html" not in names and names.index("ref/Other.html") > 1


def test_avif_behind_a_png_name_is_converted(tmp_path):
    from io import BytesIO
    from PIL import features
    from pdf2key.html import _keynote_ready
    if not features.check("avif"):
        pytest.skip("Pillow built without AVIF")
    buf = BytesIO()
    Image.new("RGB", (40, 30), (10, 200, 30)).save(buf, "AVIF")
    data = buf.getvalue()
    out = _keynote_ready(Image.open(BytesIO(data)), data)
    assert Image.open(BytesIO(out)).format in ("JPEG", "PNG")
    png = BytesIO()
    Image.new("RGB", (4, 4)).save(png, "PNG")
    assert _keynote_ready(Image.open(BytesIO(png.getvalue())), png.getvalue()) == png.getvalue()
